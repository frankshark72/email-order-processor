#!/usr/bin/env python3
"""
Importa listini Excel in EspoCRM per tutti i fornitori.
Supporta: Elmo, Prospecta, 4Power, RIB, Ermes

Prerequisiti:
  pip install requests openpyxl

Uso su Windows CMD:
  set ESPOCRM_URL=http://100.79.250.23:8080
  set ESPOCRM_API_KEY=42f83e62ee977187aa76d2e701bb6fcc
  python templates\importa_listino.py elmo     "C:\percorso\listino_elmo.xlsx"
  python templates\importa_listino.py prospecta "C:\percorso\listino_prospecta.xlsx"
  python templates\importa_listino.py 4power    "C:\percorso\listino_4power.xlsx"
  python templates\importa_listino.py rib       "C:\percorso\listino_rib.xlsx"
  python templates\importa_listino.py ermes     "C:\percorso\listino_ermes.xlsx"
"""

import sys
import os
import re
import time
import requests
import openpyxl

ESPOCRM_URL = os.environ.get("ESPOCRM_URL", "http://100.79.250.23:8080").rstrip("/")
ESPOCRM_API_KEY = os.environ.get("ESPOCRM_API_KEY", "42f83e62ee977187aa76d2e701bb6fcc")
API_BASE = f"{ESPOCRM_URL}/api/v1"
HEADERS = {"X-Api-Key": ESPOCRM_API_KEY, "Content-Type": "application/json"}

# Fogli per ogni fornitore → categoria
FOGLI_ELMO = {
    "intrusione":        "Intrusione",
    "antincendio":       "Antincendio",
    "tvcc":              "TVCC",
    "controllo accessi": "Controllo Accessi",
}


# ─────────────────────────────────────────────────────────────────────────────
# Utilities
# ─────────────────────────────────────────────────────────────────────────────

def trova_fornitore(nome_cerca: str) -> str | None:
    r = requests.get(f"{API_BASE}/Account", headers=HEADERS, params={
        "maxSize": 5, "select": "id,name",
        "where[0][type]": "contains",
        "where[0][attribute]": "name",
        "where[0][value]": nome_cerca,
    }, timeout=10)
    if not r.ok:
        return None
    lst = r.json().get("list", [])
    if lst:
        print(f"   Fornitore: {lst[0]['name']} (id: {lst[0]['id']})")
        return lst[0]["id"]
    return None


def cerca_prodotto_esistente(codice: str) -> str | None:
    r = requests.get(f"{API_BASE}/Prodotto", headers=HEADERS, params={
        "maxSize": 1, "select": "id",
        "where[0][type]": "equals",
        "where[0][attribute]": "codice",
        "where[0][value]": codice,
    }, timeout=10)
    if not r.ok or r.status_code == 404:
        return None
    lst = r.json().get("list", [])
    return lst[0]["id"] if lst else None


def pulisci_prezzo(val) -> float | None:
    if val is None:
        return None
    s = str(val).strip().replace("€", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def salva_prodotto(payload: dict, aggiorna: bool = True) -> str:
    """Inserisce o aggiorna prodotto. Ritorna 'inserito'/'aggiornato'/'saltato'/'errore'."""
    codice = payload.get("codice", "")
    existing_id = cerca_prodotto_esistente(codice) if codice else None

    if existing_id and not aggiorna:
        return "saltato"

    if existing_id:
        r = requests.patch(f"{API_BASE}/Prodotto/{existing_id}",
                           headers=HEADERS, json=payload, timeout=10)
        if r.ok:
            return "aggiornato"
        print(f"   ❌ PATCH {codice}: {r.status_code} {r.text[:100]}")
        return "errore"
    else:
        r = requests.post(f"{API_BASE}/Prodotto",
                          headers=HEADERS, json=payload, timeout=10)
        if r.ok:
            return "inserito"
        print(f"   ❌ POST {codice}: {r.status_code} {r.text[:100]}")
        return "errore"


def stampa_progresso(contatori: dict, i: int):
    if i > 0 and i % 20 == 0:
        ins = contatori["inserito"]
        agg = contatori["aggiornato"]
        print(f"      ... {ins + agg} prodotti salvati")
        time.sleep(0.5)


# ─────────────────────────────────────────────────────────────────────────────
# Parser ELMO
# ─────────────────────────────────────────────────────────────────────────────

def importa_elmo(wb: openpyxl.Workbook, fornitore_id: str) -> dict:
    contatori = {"inserito": 0, "aggiornato": 0, "saltato": 0, "errore": 0}

    for sheet_name in wb.sheetnames:
        # Trova categoria dal nome foglio
        categoria = None
        for chiave, cat in FOGLI_ELMO.items():
            if chiave in sheet_name.lower():
                categoria = cat
                break
        if not categoria:
            print(f"\n   ⏭  Foglio '{sheet_name}' non riconosciuto, salto.")
            continue

        ws = wb[sheet_name]
        print(f"\n   Foglio: {sheet_name} → {categoria}")

        # Trova header row
        col_map = {}
        header_row = None
        for row in ws.iter_rows(max_row=15):
            vals = [str(c.value or "").strip() for c in row]
            if "Prodotto" in vals or "Descrizione breve" in vals:
                header_row = row[0].row
                for cell in row:
                    v = str(cell.value or "").strip()
                    if v:
                        col_map[v] = cell.column - 1
                break

        if not header_row:
            print(f"   ⚠️  Header non trovato")
            continue

        for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
            if not any(row):
                continue

            def get(col):
                idx = col_map.get(col)
                return row[idx] if idx is not None and idx < len(row) else None

            codice = str(get("Prodotto") or "").strip()
            descrizione = str(get("Descrizione breve") or "").strip()
            if not codice or not descrizione:
                continue

            prezzo_raw = get("Prezzo")
            prezzo_su_richiesta = str(prezzo_raw or "").lower().startswith("prezzo su")
            prezzo = None if prezzo_su_richiesta else pulisci_prezzo(prezzo_raw)

            payload = {
                "name":                descrizione,
                "codice":              codice,
                "codiceAlfanumerico":  str(get("Alfanumerico") or "").strip(),
                "categoria":           categoria,
                "disponibilita":       str(get("Disponibilità") or "").strip(),
                "codiceRiparazione":   str(get("CR") or "").strip(),
                "garanzia":            str(get("G") or "").strip(),
                "prezzoSuRichiesta":   prezzo_su_richiesta,
                "unitaMisura":         "pz",
                "attivo":              True,
                "fornitoreId":         fornitore_id,
            }
            if prezzo is not None:
                payload["prezzoListino"] = prezzo

            risultato = salva_prodotto(payload)
            contatori[risultato] += 1
            stampa_progresso(contatori, i)

    return contatori


# ─────────────────────────────────────────────────────────────────────────────
# Parser PROSPECTA (cavi, prezzo al km)
# ─────────────────────────────────────────────────────────────────────────────

def importa_prospecta(wb: openpyxl.Workbook, fornitore_id: str) -> dict:
    contatori = {"inserito": 0, "aggiornato": 0, "saltato": 0, "errore": 0}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n   Foglio: {sheet_name}")

        # Trova header row (contiene "Articolo" o "Descrizione")
        col_map = {}
        header_row = None
        conf_cols = []  # colonne per conf gestite (100, 305, 500, 1000)

        for row in ws.iter_rows(max_row=15):
            vals = [str(c.value or "").strip() for c in row]
            if "Articolo" in vals or "Descrizione" in vals:
                header_row = row[0].row
                for cell in row:
                    v = str(cell.value or "").strip()
                    if v:
                        col_map[v] = cell.column - 1
                    # Colonne numeriche = conf gestite (100, 305, 500, 1000)
                    if v.isdigit():
                        conf_cols.append(v)
                break

        if not header_row:
            print(f"   ⚠️  Header non trovato")
            continue

        categoria_corrente = sheet_name

        for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
            if not any(row):
                continue

            def get(col):
                idx = col_map.get(col)
                return row[idx] if idx is not None and idx < len(row) else None

            codice = str(get("Articolo") or "").strip()

            # Riga categoria (codice vuoto, descrizione = nome categoria)
            if not codice:
                desc = str(get("Descrizione") or "").strip()
                if desc:
                    categoria_corrente = desc
                continue

            descrizione = str(get("Descrizione") or "").strip()
            if not descrizione:
                continue

            # Conf gestite: raccoglie colonne con taglie presenti
            conf_presenti = []
            for c in conf_cols:
                idx = col_map.get(c)
                if idx is not None and idx < len(row) and row[idx]:
                    conf_presenti.append(c)

            listino_col = next((k for k in col_map if "listino" in k.lower() or "prezzo" in k.lower()), None)
            prezzo = pulisci_prezzo(get(listino_col)) if listino_col else None

            payload = {
                "name":        descrizione[:255],
                "codice":      codice,
                "categoria":   categoria_corrente,
                "classeCPR":   str(get("Classe CPR") or "").strip(),
                "dop":         str(get("DoP") or "").strip(),
                "confGestite": ",".join(conf_presenti) if conf_presenti else "",
                "unitaMisura": "km",
                "attivo":      True,
                "fornitoreId": fornitore_id,
            }
            if prezzo is not None:
                payload["prezzoListino"] = prezzo

            risultato = salva_prodotto(payload)
            contatori[risultato] += 1
            stampa_progresso(contatori, i)

    return contatori


# ─────────────────────────────────────────────────────────────────────────────
# Parser generico: 4Power, RIB, Ermes
# (struttura: CODICE | DESCRIZIONE | [UM] | PREZZO)
# Le righe senza codice sono categorie
# ─────────────────────────────────────────────────────────────────────────────

def importa_generico(wb: openpyxl.Workbook, fornitore_id: str,
                     nome_fornitore: str, um_default: str = "pz") -> dict:
    contatori = {"inserito": 0, "aggiornato": 0, "saltato": 0, "errore": 0}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n   Foglio: {sheet_name}")

        # Trova header row
        col_map = {}
        header_row = None
        for row in ws.iter_rows(max_row=15):
            vals = [str(c.value or "").strip().upper() for c in row]
            if "CODICE" in vals or "ARTICOLO" in vals:
                header_row = row[0].row
                for cell in row:
                    v = str(cell.value or "").strip().upper()
                    if v:
                        col_map[v] = cell.column - 1
                break

        if not header_row:
            print(f"   ⚠️  Header non trovato")
            continue

        categoria_corrente = sheet_name

        for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
            if not any(row):
                continue

            def get(col):
                idx = col_map.get(col)
                return row[idx] if idx is not None and idx < len(row) else None

            codice = str(get("CODICE") or get("ARTICOLO") or "").strip()

            # Riga categoria (codice vuoto)
            if not codice:
                desc = str(get("DESCRIZIONE") or "").strip()
                if desc and len(desc) < 100:
                    categoria_corrente = desc
                continue

            descrizione = str(get("DESCRIZIONE") or "").strip()
            if not descrizione:
                continue

            # Unità di misura (colonna UM se presente)
            um_raw = str(get("UM") or get("UM.") or "").strip()
            um = um_raw if um_raw in ["pz", "mt", "km", "conf", "N"] else um_default

            # Prezzo: cerca colonna con "listino" o "prezzo" nel nome
            prezzo = None
            for col_name, idx in col_map.items():
                if any(k in col_name.lower() for k in ["listino", "prezzo", "euro"]):
                    prezzo = pulisci_prezzo(row[idx] if idx < len(row) else None)
                    if prezzo is not None:
                        break

            # Descrizione lunga → name troncato + descrizioneEstesa
            name_breve = descrizione[:255]
            desc_estesa = descrizione if len(descrizione) > 255 else ""

            payload = {
                "name":               name_breve,
                "codice":             codice,
                "categoria":          categoria_corrente,
                "descrizioneEstesa":  desc_estesa,
                "unitaMisura":        um,
                "attivo":             True,
                "fornitoreId":        fornitore_id,
            }
            if prezzo is not None:
                payload["prezzoListino"] = prezzo

            risultato = salva_prodotto(payload)
            contatori[risultato] += 1
            stampa_progresso(contatori, i)

    return contatori


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

FORNITORI = {
    "elmo":      ("Elmo",      "elmo"),
    "prospecta": ("Prospecta", "prospecta"),
    "4power":    ("4Power",    "4power"),
    "rib":       ("RIB",       "rib"),
    "ermes":     ("Ermes",     "ermes"),
}


def main():
    if len(sys.argv) < 3:
        print("Uso: python templates\\importa_listino.py <fornitore> <file.xlsx>")
        print("     fornitore: elmo | prospecta | 4power | rib | ermes")
        sys.exit(1)

    fornitore_key = sys.argv[1].lower()
    file_xlsx = sys.argv[2]

    if fornitore_key not in FORNITORI:
        print(f"❌ Fornitore '{fornitore_key}' non riconosciuto.")
        print(f"   Valori validi: {', '.join(FORNITORI.keys())}")
        sys.exit(1)

    nome_display, nome_cerca = FORNITORI[fornitore_key]

    print(f"\nFornitore: {nome_display}")
    print(f"File: {file_xlsx}")

    try:
        wb = openpyxl.load_workbook(file_xlsx, data_only=True)
    except Exception as e:
        print(f"❌ Errore apertura file: {e}")
        sys.exit(1)

    print(f"Fogli: {wb.sheetnames}")

    print(f"\nCerco '{nome_display}' in EspoCRM...")
    fornitore_id = trova_fornitore(nome_cerca)
    if not fornitore_id:
        print(f"❌ Fornitore '{nome_display}' non trovato in EspoCRM.")
        sys.exit(1)

    print("\nAvvio import...")
    if fornitore_key == "elmo":
        contatori = importa_elmo(wb, fornitore_id)
    elif fornitore_key == "prospecta":
        contatori = importa_prospecta(wb, fornitore_id)
    else:
        contatori = importa_generico(wb, fornitore_id, nome_display)

    print(f"\n{'='*50}")
    print(f"TOTALE {nome_display}:")
    print(f"  ✅ Inseriti:   {contatori['inserito']}")
    print(f"  🔄 Aggiornati: {contatori['aggiornato']}")
    print(f"  ⏭  Saltati:    {contatori['saltato']}")
    print(f"  ❌ Errori:     {contatori['errore']}")
    print("="*50)


if __name__ == "__main__":
    main()
